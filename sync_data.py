import os
import sys
import glob
import json
import re
from datetime import datetime
import openpyxl

# Set output encoding to UTF-8
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def find_excel_file():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    parent_dir = os.path.dirname(current_dir)
    
    # Priority paths
    candidates = []
    candidates.extend(glob.glob(os.path.join(current_dir, "*จิตวิทยา*.xlsx")))
    candidates.extend(glob.glob(os.path.join(current_dir, "*TCAS*.xlsx")))
    candidates.extend(glob.glob(os.path.join(parent_dir, "*จิตวิทยา*.xlsx")))
    candidates.extend(glob.glob(os.path.join(parent_dir, "*TCAS*.xlsx")))
    
    seen = set()
    valid_candidates = []
    for c in candidates:
        if c not in seen and os.path.exists(c):
            seen.add(c)
            valid_candidates.append(c)
            
    if not valid_candidates:
        fallback = r"E:\OneDrive\Projects\Other Task\สรุป_TCAS70_Port_จิตวิทยา.xlsx"
        if os.path.exists(fallback):
            return fallback
        return None
    return valid_candidates[0]

def parse_excel_to_json(excel_path):
    print(f"Loading Excel file: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active
    
    headers = []
    for col in range(1, sheet.max_column + 1):
        val = sheet.cell(1, col).value
        headers.append(str(val).strip() if val is not None else f"col_{col}")
        
    print(f"Headers found: {headers}")
    
    universities = []
    
    brand_map = {
        "มหาวิทยาลัยศรีนครินทรวิโรฒ": {
            "shortName": "มศว (SWU)",
            "abbr": "SWU",
            "theme": "swu",
            "badgeColor": "bg-rose-100 text-rose-800 border-rose-300 dark:bg-rose-900/70 dark:text-rose-200 dark:border-rose-700",
            "accentColor": "from-rose-500 to-red-600",
            "tagColor": "bg-rose-100 text-rose-800 dark:bg-rose-950 dark:text-rose-200"
        },
        "มหาวิทยาลัยเกษตรศาสตร์": {
            "shortName": "มก. (KU)",
            "abbr": "KU",
            "theme": "ku",
            "badgeColor": "bg-emerald-100 text-emerald-800 border-emerald-300 dark:bg-emerald-900/70 dark:text-emerald-200 dark:border-emerald-700",
            "accentColor": "from-emerald-500 to-teal-600",
            "tagColor": "bg-emerald-100 text-emerald-800 dark:bg-emerald-950 dark:text-emerald-200"
        },
        "มหาวิทยาลัยธรรมศาสตร์": {
            "shortName": "มธ. (TU)",
            "abbr": "TU",
            "theme": "tu",
            "badgeColor": "bg-amber-100 text-amber-900 border-amber-300 dark:bg-amber-900/70 dark:text-amber-200 dark:border-amber-700",
            "accentColor": "from-amber-500 to-red-500",
            "tagColor": "bg-amber-100 text-amber-900 dark:bg-amber-950 dark:text-amber-200"
        },
        "มหาวิทยาลัยศิลปากร": {
            "shortName": "มศก. (SU)",
            "abbr": "SU",
            "theme": "su",
            "badgeColor": "bg-blue-100 text-blue-800 border-blue-300 dark:bg-blue-900/70 dark:text-blue-200 dark:border-blue-700",
            "accentColor": "from-blue-500 to-indigo-600",
            "tagColor": "bg-blue-100 text-blue-800 dark:bg-blue-950 dark:text-blue-200"
        }
    }
    
    for row_idx in range(2, sheet.max_row + 1):
        row_vals = {}
        for col_idx, h in enumerate(headers):
            cell_val = sheet.cell(row_idx, col_idx + 1).value
            row_vals[h] = str(cell_val).strip() if cell_val is not None else ""
            
        uni_name = row_vals.get("มหาวิทยาลัย", "")
        if not uni_name:
            continue
            
        guidelines = row_vals.get("แนวทางการทำ Port", "")
        gpax_match = re.search(r"GPAX.*?(\d+\.\d+)", guidelines, re.IGNORECASE)
        gpax_min = gpax_match.group(1) if gpax_match else None
        
        curriculum = row_vals.get("หลักสูตรเป็นแบบไหน", "")
        degree_type = "ไม่ระบุ"
        if "วท.บ." in curriculum:
            degree_type = "วท.บ."
        elif "ศศ.บ." in curriculum:
            degree_type = "ศศ.บ."
            
        tracks = []
        quota_text = row_vals.get("จำนวนรับรอบ Port", "")
        combined_text = curriculum + " " + quota_text
        if "คลินิก" in combined_text:
            tracks.append("จิตวิทยาคลินิก")
        if "ชุมชน" in combined_text:
            tracks.append("จิตวิทยาชุมชน")
        if "พัฒนาการ" in combined_text:
            tracks.append("จิตวิทยาพัฒนาการ")
        if "อุตสาหกรรม" in combined_text:
            tracks.append("จิตวิทยาอุตสาหกรรมและองค์การ")
        if "ธุรกิจ" in combined_text:
            tracks.append("จิตวิทยาธุรกิจ")
        if not tracks and "จิตวิทยา" in combined_text:
            tracks.append("จิตวิทยาทั่วไป")
            
        brand = brand_map.get(uni_name, {
            "shortName": uni_name[:10],
            "abbr": "UNI",
            "theme": "slate",
            "badgeColor": "bg-slate-100 text-slate-800 border-slate-300 dark:bg-slate-700 dark:text-slate-200 dark:border-slate-600",
            "accentColor": "from-slate-600 to-slate-800",
            "tagColor": "bg-slate-100 text-slate-800 dark:bg-slate-700 dark:text-slate-200"
        })
        
        uni_data = {
            "id": f"uni-{row_idx-1}",
            "university": uni_name,
            "shortName": brand["shortName"],
            "abbr": brand["abbr"],
            "theme": brand["theme"],
            "badgeColor": brand["badgeColor"],
            "accentColor": brand["accentColor"],
            "tagColor": brand["tagColor"],
            "faculty": row_vals.get("ชื่อคณะ", ""),
            "major": row_vals.get("ชื่อสาขาวิชา", "จิตวิทยา"),
            "degreeType": degree_type,
            "tracks": tracks,
            "portDate": row_vals.get("วันรับสมัคร Port", ""),
            "quotaCount": row_vals.get("จำนวนรับรอบ Port", ""),
            "quotaRoundDate": row_vals.get("รอบโควต้า", ""),
            "portGuidelines": guidelines,
            "website": row_vals.get("Website ที่เกี่ยวข้อง", ""),
            "curriculum": curriculum,
            "gpaxMin": gpax_min
        }
        universities.append(uni_data)
        
    output_data = {
        "metadata": {
            "title": "สรุป TCAS 70 รอบ Portfolio สาขาจิตวิทยา",
            "description": "ข้อมูลเกณฑ์การรับสมัคร จำนวนรับ วันเปิดรับ และแนวทางการทำพอร์ตจิตวิทยา",
            "sourceExcel": os.path.basename(excel_path),
            "lastSyncedAt": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            "totalUniversities": len(universities)
        },
        "universities": universities
    }
    
    return output_data

def main():
    excel_file = find_excel_file()
    if not excel_file:
        print("Error: ไม่พบไฟล์ Excel 'สรุป_TCAS70_Port_จิตวิทยา.xlsx'")
        sys.exit(1)
        
    data = parse_excel_to_json(excel_file)
    
    script_dir = os.path.dirname(os.path.abspath(__file__))
    output_dir = os.path.join(script_dir, "src", "data")
    os.makedirs(output_dir, exist_ok=True)
    
    output_file = os.path.join(output_dir, "tcas_data.json")
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        
    print(f"Sync complete! Successfully written {len(data['universities'])} universities to: {output_file}")

if __name__ == "__main__":
    main()
